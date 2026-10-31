"""산림임업용어사전.xlsx → GitHub에서 바로 읽히는 초성별 md 파일과 목차 README를 만든다.

입력: 07_산림임업용어사전/산림임업용어사전.xlsx ('산림임업용어사전' 시트, '정보' 시트)
출력: 07_산림임업용어사전/초성별/NN_초성.md (GitHub 렌더링 한도를 넘지 않게 큰 초성은 나눔), README.md
여러 번 실행해도 결과가 같다. 실행: python3 tools/build_glossary_md.py  (openpyxl 필요)
"""
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DIR = ROOT / "07_산림임업용어사전"
OUT = DIR / "초성별"
MAX_BYTES = 450_000  # GitHub md 미리보기가 잘리지 않을 크기

CHO = "ㄱㄲㄴㄷㄸㄹㅁㅂㅃㅅㅆㅇㅈㅉㅊㅋㅌㅍㅎ"
MERGE = {"ㄲ": "ㄱ", "ㄸ": "ㄷ", "ㅃ": "ㅂ", "ㅆ": "ㅅ", "ㅉ": "ㅈ"}
GROUPS = list("ㄱㄴㄷㄹㅁㅂㅅㅇㅈㅊㅋㅌㅍㅎ") + ["기타"]


def initial(word):
    c = word[:1]
    if "가" <= c <= "힣":
        ch = CHO[(ord(c) - 0xAC00) // 588]
        return MERGE.get(ch, ch)
    return "기타"


def esc(s):
    return str(s or "").replace("\r", "").strip()


def entry(row):
    no, ko, en, hanja, desc, _att, link = row
    meta = " · ".join(x for x in (esc(hanja), esc(en)) if x)
    head = f"### {esc(ko)}" + (f" <sub>{meta}</sub>" if meta else "")
    body = esc(desc).replace("\n", "  \n") or "_(설명 없음)_"
    tail = f"<sub>No.{no} · [산림청 원문]({link})</sub>" if link else f"<sub>No.{no}</sub>"
    return f"{head}\n\n{body}\n\n{tail}\n"


def chunks(entries):
    part, size = [], 0
    for e in entries:
        b = len(e.encode())
        if part and size + b > MAX_BYTES:
            yield part
            part, size = [], 0
        part.append(e)
        size += b
    if part:
        yield part


def main():
    wb = openpyxl.load_workbook(DIR / "산림임업용어사전.xlsx", read_only=True)
    info = {r[0]: r[1] for r in wb["정보"].iter_rows(values_only=True) if r[0]}
    rows = [r for r in wb["산림임업용어사전"].iter_rows(min_row=2, values_only=True) if r[1]]

    groups = {g: [] for g in GROUPS}
    for r in rows:
        groups[initial(esc(r[1]))].append(r)

    OUT.mkdir(exist_ok=True)
    for old in OUT.glob("*.md"):
        old.unlink()

    toc = []
    for i, g in enumerate(GROUPS, 1):
        items = groups[g]
        if not items:
            continue
        parts = list(chunks([entry(r) for r in items]))
        links = []
        for k, part in enumerate(parts, 1):
            suffix = f"-{k}" if len(parts) > 1 else ""
            name = f"{i:02d}_{g}{suffix}.md"
            first, last = part[0].split("\n")[0][4:].split(" <sub>")[0], part[-1].split("\n")[0][4:].split(" <sub>")[0]
            title = f"{g}{f' ({k}/{len(parts)})' if len(parts) > 1 else ''}"
            (OUT / name).write_text(
                f"# 산림임업용어사전 — {title}\n\n"
                f"[← 목차](../README.md) · {len(part):,}개 용어 · {first} ~ {last}\n\n---\n\n"
                + "\n".join(part),
                encoding="utf-8",
            )
            links.append(f"[{title}](초성별/{name.replace(' ', '%20')}) <sub>{first} ~ {last}</sub>")
        toc.append(f"| **{g}** | {len(items):,} | {' · '.join(links)} |")

    (DIR / "README.md").write_text(
        "# 산림임업용어사전\n\n"
        f"> 출처: {info.get('출처', '산림청')} ({info.get('URL', '')})  \n"
        f"> 수집일 {info.get('수집일', '')} · 총 **{len(rows):,}개** 용어 (국문명·한자·영문·설명·원문 링크)\n\n"
        "## 보는 방법\n\n"
        "| 파일 | 용도 |\n|---|---|\n"
        "| 아래 **초성별 목차** | GitHub에서 바로 읽기 — 초성을 누르고 `Ctrl/Cmd + F`로 찾기 |\n"
        "| `산림임업용어사전.html` | 내려받아 더블클릭 → 검색·초성 탭이 있는 사전 앱 (오프라인) |\n"
        "| `산림임업용어사전.xlsx` | 엑셀 원본 데이터 (필터·정렬용) |\n\n"
        "## 초성별 목차\n\n| 초성 | 용어 수 | 파일 |\n|---|---:|---|\n"
        + "\n".join(toc)
        + "\n\n---\n\n엑셀을 고쳤으면 `python3 tools/build_glossary_md.py`로 이 목차와 초성별 파일을 다시 만드세요.\n",
        encoding="utf-8",
    )
    print(f"{len(rows)} terms → {len(list(OUT.glob('*.md')))} files")


if __name__ == "__main__":
    main()
